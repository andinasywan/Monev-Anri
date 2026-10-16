"""Extract the public columns of the ANRI Monev Renja workbook into data.js + meta.json.

Usage: python3 extract.py <workbook.xlsx> <out_dir> <source_modified_iso>

Only columns found by their header label are read (an allowlist). The internal
columns (O, Q, R, S, T, U in the current layout) are never read or written.
"""
import json
import re
import sys
from datetime import datetime, timezone

import openpyxl
from openpyxl.utils import get_column_letter

MONTHS = ["JANUARI", "FEBRUARI", "MARET", "APRIL", "MEI", "JUNI", "JULI",
          "AGUSTUS", "SEPTEMBER", "OKTOBER", "NOVEMBER", "DESEMBER"]
SHEET_RE = re.compile(r"^\s*(" + "|".join(MONTHS) + r")\s+(\d{4})\s*$", re.I)


def norm(v):
    return " ".join(str(v).upper().replace("\n", " ").split()) if isinstance(v, str) else ""


# (field, header row, matcher). First match from the left wins.
SD = r"S\.?\s?D\.?"
RULES = [
    ("pagu_dipa", 5, lambda s: s.startswith("PAGU DIPA")),
    ("pagu_aktif", 5, lambda s: s.startswith("PAGU AKTIF")),
    ("sp2d_bln", 5, lambda s: re.match(r"^REALISASI SP2D (?!" + SD + ")", s)),
    ("spm_bln", 5, lambda s: re.match(r"^REALISASI SPM (?!" + SD + ")", s)),
    ("sp2d_sd", 5, lambda s: re.match(r"^REALISASI SP2D " + SD, s)),
    ("spm_sd", 5, lambda s: re.match(r"^REALISASI SPM " + SD, s)),
    ("pct_sp2d", 5, lambda s: s.startswith("% REALISASI SP2D")),
    ("pct_spm", 5, lambda s: s.startswith("% REALISASI SPM")),
    ("target", 5, lambda s: s == "TARGET"),
    ("satuan", 5, lambda s: s == "SATUAN"),
    ("fisik_bln", 5, lambda s: re.match(r"^REALISASI FISIK (?!" + SD + ")", s)),
    ("fisik_sd", 5, lambda s: re.match(r"^REALISASI FISIK " + SD, s)),
    ("pct_fisik", 5, lambda s: s.startswith("% KINERJA FISIK")),
    ("catatan", 4, lambda s: s == "CATATAN"),
]
FIELDS = ["baris", "kode", "uraian", "level", "kelompok", "unit"] + [r[0] for r in RULES]
NUMERIC = {"pagu_dipa", "pagu_aktif", "sp2d_bln", "spm_bln", "sp2d_sd", "spm_sd", "pct_sp2d",
           "pct_spm", "target", "fisik_bln", "fisik_sd", "pct_fisik"}


# Internal-only columns: never read, in any sheet, whatever their header says.
BLOCKED = {"O", "Q", "R", "S", "T", "U"}


def locate(ws):
    cols = {"kode": 1, "uraian": 2}
    for field, row, match in RULES:
        for c in range(3, ws.max_column + 1):
            if get_column_letter(c) in BLOCKED:
                continue
            if match(norm(ws.cell(row, c).value)):
                cols[field] = c
                break
    return cols


def code_text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def num(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 10) if not float(v).is_integer() else int(v)
    s = str(v).strip().replace("%", "")
    if not s or s == "-":
        return None
    try:
        return float(s.replace(".", "").replace(",", ".")) if "," in s else float(s)
    except ValueError:
        return None


def level_of(kode, uraian, target):
    if isinstance(target, (int, float)):
        return "ro"
    if kode == "087":
        return "total"
    if re.fullmatch(r"\d{3}\.[A-Z]{2}", kode):
        return "program"
    if re.fullmatch(r"\d{6}", kode):
        return "satker"
    if re.fullmatch(r"\d{4}", kode):
        return "kegiatan"
    if re.fullmatch(r"[A-Z][A-Z0-9]{2}", kode):
        return "kro"
    if not kode:
        u = uraian.upper()
        if u.startswith("ESELON II MANDIRI"):
            return "kelompok"
        return "eselon1" if u.startswith(("DEPUTI", "SEKRETARIAT UTAMA")) else "unit"
    return "lainnya"


def extract_sheet(ws):
    cols = locate(ws)
    rows = []
    kelompok = ""
    unit = ""
    eselon1 = ""
    for r in range(7, ws.max_row + 1):
        uraian = ws.cell(r, cols["uraian"]).value
        uraian = " ".join(str(uraian).split()) if uraian not in (None, "") else ""
        if not uraian:
            continue
        kode = code_text(ws.cell(r, cols["kode"]).value)
        rec = {"baris": r, "kode": kode, "uraian": uraian}
        for f, c in cols.items():
            if f in ("kode", "uraian"):
                continue
            v = ws.cell(r, c).value
            rec[f] = num(v) if f in NUMERIC else (" ".join(str(v).split()) if v not in (None, "") else "")
        lvl = level_of(kode, uraian, rec.get("target"))
        rec["level"] = lvl
        if lvl == "eselon1":
            eselon1 = uraian
            kelompok = uraian
            unit = ""
        elif lvl == "unit":
            # Direktorat sit under a Deputi, Biro under Sekretariat Utama;
            # Pusat, Balai and Inspektorat stand on their own.
            up = uraian.upper()
            under = (up.startswith("DIREKTORAT") and eselon1.upper().startswith("DEPUTI")) or \
                    (up.startswith("BIRO") and eselon1.upper().startswith("SEKRETARIAT"))
            kelompok = eselon1 if under else "Pusat, Balai dan Inspektorat"
            unit = uraian
        elif lvl in ("total", "program", "satker", "kelompok"):
            eselon1 = ""
            kelompok = ""
            unit = ""
        rec["kelompok"] = kelompok if lvl not in ("total", "program", "satker") else ""
        rec["unit"] = unit if lvl not in ("total", "program", "satker", "eselon1") else ""
        rows.append([rec.get(f) for f in FIELDS])
    return cols, rows


def main():
    src, out_dir, modified = sys.argv[1], sys.argv[2], sys.argv[3]
    wb = openpyxl.load_workbook(src, data_only=True, read_only=False)
    months = []
    data = {}
    picked = {}
    for ws in wb.worksheets:
        m = SHEET_RE.match(ws.title)
        if not m:
            continue
        cols, rows = extract_sheet(ws)
        if not any(r[FIELDS.index("level")] == "total" for r in rows):
            continue
        no = MONTHS.index(m.group(1).upper()) + 1
        months.append({"sheet": ws.title.strip(), "bulan": m.group(1).capitalize(), "no": no,
                       "tahun": int(m.group(2))})
        data[ws.title.strip()] = rows
        picked[ws.title.strip()] = "".join(sorted({get_column_letter(c) for c in cols.values()},
                                                  key=lambda x: (len(x), x)))
    months.sort(key=lambda x: (x["tahun"], x["no"]))
    payload = {
        "generated": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "sourceModified": modified,
        "fields": FIELDS,
        "months": months,
        "rows": data,
    }
    js = "window.MONEV=" + json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + ";\n"
    with open(f"{out_dir}/data.js", "w", encoding="utf-8") as f:
        f.write(js)
    with open(f"{out_dir}/meta.json", "w", encoding="utf-8") as f:
        json.dump({"sourceModified": modified, "generated": payload["generated"],
                   "months": [m["sheet"] for m in months]}, f)
    print("months:", [m["sheet"] for m in months])
    print("columns read per sheet:", picked)
    print("data.js bytes:", len(js.encode("utf-8")))


if __name__ == "__main__":
    main()
